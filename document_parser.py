from __future__ import annotations

from pathlib import Path
import re
from collections.abc import Callable

import ocr_client


TEXT_SUFFIXES = {".md", ".markdown", ".txt"}
LONG_DOCUMENT_PAGE_THRESHOLD = 8
LONG_DOCUMENT_CHAR_THRESHOLD = 12000
NOISE_LINE_PATTERNS = (
    r"^\d+\s*/\s*\d+$",
    r"^page\s+\d+(\s+of\s+\d+)?$",
    r"^第\s*\d+\s*页\s*(共\s*\d+\s*页)?$",
    r"^-+\s*\d+\s*-+$",
    r"^confidential$",
    r"^copyright\s+.*$",
    r"^all rights reserved\.?$",
)


def read_text_file(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030", "gbk"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="ignore")


def read_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("缺少 python-docx 依赖，请运行 pip install -r requirements.txt") from exc

    document = Document(path)
    blocks: list[str] = []
    blocks.extend(paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip())
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return "\n".join(blocks).strip()


VisionRecognizer = Callable[[list[Path]], str]


def _ocr_pdf_image(image_path: Path, visual_recognizer: VisionRecognizer | None = None, prefer_visual: bool = False) -> tuple[str, str]:
    if prefer_visual and visual_recognizer:
        return visual_recognizer([image_path]).strip(), "ai_vision"
    try:
        text = ocr_client.recognize_screenshots([image_path]).strip()
        if text and "No usable text recognized." not in text:
            return text, "local_ocr"
    except Exception:
        if not visual_recognizer:
            raise
    if visual_recognizer:
        return visual_recognizer([image_path]).strip(), "ai_vision"
    return text.strip(), "local_ocr"


def read_pdf(path: Path, visual_recognizer: VisionRecognizer | None = None, prefer_visual: bool = False) -> str:
    try:
        import fitz
    except ImportError as exc:
        raise RuntimeError("缺少 PyMuPDF 依赖，请运行 pip install -r requirements.txt") from exc

    doc = fitz.open(path)
    blocks: list[str] = []
    ocr_images: list[Path] = []
    render_dir = ocr_client.CACHE_DIR / "pdf_pages"
    render_dir.mkdir(parents=True, exist_ok=True)
    try:
        for index, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            if len(text) >= 30:
                blocks.append(f"[PDF Page {index}]\n{text}")
                continue
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            image_path = render_dir / f"{path.stem}_{index}.png"
            pixmap.save(image_path)
            ocr_images.append(image_path)
        for image_path in ocr_images:
            page_number = image_path.stem.rsplit("_", 1)[-1]
            text, mode = _ocr_pdf_image(image_path, visual_recognizer=visual_recognizer, prefer_visual=prefer_visual)
            if text:
                blocks.append(f"[PDF Page {page_number} {mode}]\n{text}")
    finally:
        doc.close()
    return clean_extracted_document_text("\n\n---\n\n".join(block for block in blocks if block.strip()).strip())


def read_pdf_pages(path: Path, visual_recognizer: VisionRecognizer | None = None, prefer_visual: bool = False) -> list[dict]:
    try:
        import fitz
    except ImportError as exc:
        raise RuntimeError("缺少 PyMuPDF 依赖，请运行 pip install -r requirements.txt") from exc

    doc = fitz.open(path)
    pages: list[dict] = []
    render_dir = ocr_client.CACHE_DIR / "pdf_pages"
    render_dir.mkdir(parents=True, exist_ok=True)
    try:
        for index, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            used_ocr = False
            ocr_mode = ""
            if len(text) < 30:
                pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                image_path = render_dir / f"{path.stem}_{index}.png"
                pixmap.save(image_path)
                text, ocr_mode = _ocr_pdf_image(image_path, visual_recognizer=visual_recognizer, prefer_visual=prefer_visual)
                used_ocr = True
            cleaned_text = clean_extracted_document_text(text)
            pages.append(
                {
                    "page": index,
                    "text": cleaned_text,
                    "char_count": len(cleaned_text),
                    "used_ocr": used_ocr,
                    "ocr_mode": ocr_mode,
                }
            )
    finally:
        doc.close()
    return pages


def document_pages(path: Path, visual_recognizer: VisionRecognizer | None = None, prefer_visual: bool = False) -> list[dict]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return read_pdf_pages(path, visual_recognizer=visual_recognizer, prefer_visual=prefer_visual)
    text = extract_text(path, visual_recognizer=visual_recognizer, prefer_visual=prefer_visual)
    chunks = chunk_text(text, max_chars=4500)
    return [
        {
            "page": index,
            "text": chunk,
            "char_count": len(chunk),
            "used_ocr": False,
        }
        for index, chunk in enumerate(chunks, start=1)
    ]


def chunk_text(text: str, max_chars: int = 6000) -> list[str]:
    text = text.strip()
    if len(text) <= max_chars:
        return [text] if text else []
    paragraphs = re.split(r"\n\s*\n", text)
    chunks: list[str] = []
    current: list[str] = []
    current_len = 0
    for paragraph in paragraphs:
        paragraph = paragraph.strip()
        if not paragraph:
            continue
        if current and current_len + len(paragraph) > max_chars:
            chunks.append("\n\n".join(current))
            current = []
            current_len = 0
        if len(paragraph) > max_chars:
            for start in range(0, len(paragraph), max_chars):
                part = paragraph[start : start + max_chars]
                if part:
                    chunks.append(part)
            continue
        current.append(paragraph)
        current_len += len(paragraph)
    if current:
        chunks.append("\n\n".join(current))
    return chunks


def clean_extracted_document_text(text: str) -> str:
    lines = [line.strip() for line in text.replace("\r\n", "\n").replace("\r", "\n").splitlines()]
    repeated_short_lines = _repeated_short_lines(lines)
    cleaned: list[str] = []
    previous_blank = False
    for line in lines:
        if not line:
            if cleaned and not previous_blank:
                cleaned.append("")
            previous_blank = True
            continue
        normalized = re.sub(r"\s+", " ", line).strip()
        if _is_document_noise_line(normalized, repeated_short_lines):
            continue
        cleaned.append(normalized)
        previous_blank = False
    return _trim_document_noise_edges("\n".join(cleaned).strip())


def _repeated_short_lines(lines: list[str]) -> set[str]:
    counts: dict[str, int] = {}
    for line in lines:
        normalized = re.sub(r"\s+", " ", line).strip()
        if 0 < len(normalized) <= 80:
            counts[normalized.lower()] = counts.get(normalized.lower(), 0) + 1
    return {line for line, count in counts.items() if count >= 3}


def _is_document_noise_line(line: str, repeated_short_lines: set[str]) -> bool:
    lower = line.lower()
    if lower in repeated_short_lines:
        return True
    if _looks_like_catalog_line(lower):
        return True
    if _looks_like_contact_or_footer_line(lower):
        return True
    return any(re.match(pattern, lower, flags=re.IGNORECASE) for pattern in NOISE_LINE_PATTERNS)


def _looks_like_catalog_line(line: str) -> bool:
    if line in {"contents", "table of contents", "目录"}:
        return True
    if re.match(r"^[\divx一二三四五六七八九十a-z]+[.)、]\s+.+?\.{2,}\s*\d+$", line, flags=re.IGNORECASE):
        return True
    if re.match(r"^.+?\.{3,}\s*\d+$", line):
        return True
    if re.match(r"^(figure|table)\s+\d+[.: -].*$", line, flags=re.IGNORECASE):
        return True
    return False


def _looks_like_contact_or_footer_line(line: str) -> bool:
    if len(line) <= 120 and ("@" in line or "www." in line or "http://" in line or "https://" in line):
        return True
    if re.match(r"^(tel|phone|email|fax)[:：]\s*.+$", line, flags=re.IGNORECASE):
        return True
    return False


def _trim_document_noise_edges(text: str) -> str:
    if not text:
        return text
    lines = text.splitlines()
    while lines and _is_probably_noise_edge(lines[0]):
        lines.pop(0)
    while lines and _is_probably_noise_edge(lines[-1]):
        lines.pop()
    return "\n".join(lines).strip()


def _is_probably_noise_edge(line: str) -> bool:
    normalized = re.sub(r"\s+", " ", line).strip().lower()
    if not normalized:
        return True
    if len(normalized) <= 6 and re.fullmatch(r"[\d./-]+", normalized):
        return True
    return normalized in {"home", "login", "sign in", "contents", "table of contents", "目录"}


def pages_text(pages: list[dict], page_start: int | None = None, page_end: int | None = None) -> str:
    selected = []
    for page in pages:
        page_number = int(page.get("page") or 0)
        if page_start is not None and page_number < page_start:
            continue
        if page_end is not None and page_number > page_end:
            continue
        text = (page.get("text") or "").strip()
        if text:
            selected.append(f"[Page {page_number}]\n{text}")
    return "\n\n---\n\n".join(selected).strip()


def parse_page_ranges(value: str | None, page_count: int) -> list[tuple[int, int]]:
    text = (value or "").strip()
    if not text:
        return [(1, page_count)] if page_count else []
    ranges: list[tuple[int, int]] = []
    for part in re.split(r"[;?,\s]+", text):
        part = part.strip()
        if not part:
            continue
        match = re.fullmatch(r"(\d+)(?:\s*-\s*(\d+))?", part)
        if not match:
            raise ValueError(f"?????????{part}")
        start_page = int(match.group(1))
        end_page = int(match.group(2) or start_page)
        if start_page > end_page:
            start_page, end_page = end_page, start_page
        if page_count:
            start_page = max(1, min(start_page, page_count))
            end_page = max(1, min(end_page, page_count))
        ranges.append((start_page, end_page))
    return merge_page_ranges(ranges)


def merge_page_ranges(ranges: list[tuple[int, int]]) -> list[tuple[int, int]]:
    if not ranges:
        return []
    ordered = sorted(ranges)
    merged = [ordered[0]]
    for start_page, end_page in ordered[1:]:
        last_start, last_end = merged[-1]
        if start_page <= last_end + 1:
            merged[-1] = (last_start, max(last_end, end_page))
        else:
            merged.append((start_page, end_page))
    return merged


def filter_pages_by_ranges(pages: list[dict], ranges: list[tuple[int, int]]) -> list[dict]:
    if not ranges:
        return []
    selected = []
    for page in pages:
        page_number = int(page.get("page") or 0)
        if any(start_page <= page_number <= end_page for start_page, end_page in ranges):
            selected.append(page)
    return selected


def format_page_ranges(ranges: list[tuple[int, int]]) -> str:
    parts = []
    for start_page, end_page in ranges:
        parts.append(str(start_page) if start_page == end_page else f"{start_page}-{end_page}")
    return ";".join(parts)


def compact_pages_for_prompt(pages: list[dict], max_chars_per_page: int = 900) -> str:
    blocks = []
    for page in pages:
        text = " ".join((page.get("text") or "").split())
        if len(text) > max_chars_per_page:
            text = text[:max_chars_per_page] + "..."
        blocks.append(f"第 {page.get('page')} 页（{page.get('char_count', 0)}字）：{text}")
    return "\n".join(blocks)


def extract_text(path: Path, visual_recognizer: VisionRecognizer | None = None, prefer_visual: bool = False) -> str:
    suffix = path.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        text = read_text_file(path)
    elif suffix == ".docx":
        text = read_docx(path)
    elif suffix == ".pdf":
        text = read_pdf(path, visual_recognizer=visual_recognizer, prefer_visual=prefer_visual)
    elif suffix == ".doc":
        raise ValueError("暂不支持 .doc，请另存为 .docx 后上传")
    else:
        raise ValueError(f"不支持的文件类型：{suffix or path.name}")
    text = text.strip()
    if not text:
        raise ValueError(f"未能从文件中提取到可用文本：{path.name}")
    return clean_extracted_document_text(text)


def recognize_files(
    files: list[dict],
    root: Path,
    visual_recognizer: VisionRecognizer | None = None,
    prefer_visual: bool = False,
    storage_root: Path | None = None,
) -> str:
    blocks = []
    for index, item in enumerate(files, start=1):
        file_path = Path(str(item["file_path"]))
        if not file_path.is_absolute():
            candidates = [
                root / file_path,
                root / "data" / "runtime" / file_path,
            ]
            if storage_root is not None:
                candidates.insert(0, storage_root / file_path)
            local_appdata = Path.home() / "AppData" / "Local" / "FigureLearning"
            candidates.append(local_appdata / file_path)
            file_path = next((candidate for candidate in candidates if candidate.exists()), candidates[0])
        text = extract_text(file_path, visual_recognizer=visual_recognizer, prefer_visual=prefer_visual)
        name = item.get("original_name") or file_path.name
        blocks.append(f"[File {index}: {name}]\n{text}")
    return "\n\n---\n\n".join(blocks)
